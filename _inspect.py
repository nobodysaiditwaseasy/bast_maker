from docx import Document
doc = Document('template/template_laporan_perjadin.docx')
print(f'Paragraphs: {len(doc.paragraphs)}')
print(f'Tables: {len(doc.tables)}')
print(f'Sections: {len(doc.sections)}')
s = doc.sections[0]
print(f'Page W: {s.page_width} H: {s.page_height}')
print(f'Margins: L={s.left_margin} R={s.right_margin} T={s.top_margin} B={s.bottom_margin}')
# Check for page breaks
for i, p in enumerate(doc.paragraphs):
    xml = p._element.xml
    if '<w:br' in xml:
        print(f'BR found at P{i}: [{p.text.strip()[:50]}]')
# Check for images
rels = doc.part.rels
for r in rels.values():
    if 'image' in r.reltype:
        print(f'Image: {r.target_ref}')
# Check all paragraph text
print('\n--- All paragraphs ---')
for i, p in enumerate(doc.paragraphs):
    print(f'P{i}: [{p.text.strip()[:80]}]')
